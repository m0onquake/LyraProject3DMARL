from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ACCENT = "1F4E79"
ACCENT_LIGHT = "EAF2F8"
BORDER = "D9D9D9"
BODY_FONT = "Microsoft YaHei"
MONO_FONT = "Cascadia Mono"


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_inline(paragraph, text: str, *, bold=False, italic=False, size=10.5):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, MONO_FONT, max(9.0, size - 0.5), color="333333")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, BODY_FONT, size, bold=bold)
            run.italic = italic


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.35

    title = doc.styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)

    for style_name, size, before, after in (
        ("Heading 1", 16, 16, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.25


def add_footer(doc: Document, short_title: str):
    for section in doc.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(short_title)
        set_run_font(run, BODY_FONT, 8.5, color="666666")
        run = paragraph.add_run("  |  ")
        set_run_font(run, BODY_FONT, 8.5, color="999999")
        run = paragraph.add_run("Page ")
        set_run_font(run, BODY_FONT, 8.5, color="666666")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        paragraph._p.append(fld)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = True
    set_table_borders(table)
    for r_idx, values in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, ACCENT)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, ACCENT_LIGHT)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            value = values[c_idx] if c_idx < len(values) else ""
            add_inline(paragraph, value, bold=(r_idx == 0), size=9.2)
            if r_idx == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", part or "") for part in parts):
            rows.append(parts)
        index += 1
    return rows, index


def build_doc(source: Path, output: Path, short_title: str):
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)
    configure_styles(doc)
    add_footer(doc, short_title)

    paragraph_buffer: list[str] = []
    in_code = False
    code_buffer: list[str] = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            add_inline(paragraph, " ".join(item.strip() for item in paragraph_buffer))
            paragraph_buffer = []

    index = 0
    title_written = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if line.startswith("```"):
            flush_paragraph()
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.3)
                paragraph.paragraph_format.right_indent = Inches(0.2)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.paragraph_format.line_spacing = 1.05
                for pos, code_line in enumerate(code_buffer):
                    if pos:
                        paragraph.add_run().add_break()
                    run = paragraph.add_run(code_line)
                    set_run_font(run, MONO_FONT, 9, color="333333")
                in_code = False
            index += 1
            continue
        if in_code:
            code_buffer.append(raw)
            index += 1
            continue
        if not line:
            flush_paragraph()
            index += 1
            continue
        if line.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            flush_paragraph()
            if not title_written:
                paragraph = doc.add_paragraph(style="Title")
                add_inline(paragraph, line[2:].strip(), bold=True, size=24)
                title_written = True
            index += 1
            continue
        if line.startswith("## "):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:].strip(), bold=True, size=16)
            index += 1
            continue
        if line.startswith("### "):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:].strip(), bold=True, size=13)
            index += 1
            continue
        if re.match(r"^-\s+", line):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^-\s+", "", line))
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
            index += 1
            continue
        if line.endswith("  ") or re.match(r"^(版本|适用工程|目标地图|实施分支)：", line):
            flush_paragraph()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, line.rstrip(), size=9.5)
            index += 1
            continue
        paragraph_buffer.append(line)
        index += 1

    flush_paragraph()
    doc.core_properties.title = lines[0].removeprefix("# ").strip()
    doc.core_properties.subject = "UE5.3 Lyra TacticalMARL 红方规则对抗升级"
    doc.core_properties.author = "TacticalMARL Project"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    build_doc(
        ROOT / "红方规则对抗MARL场景需求说明.md",
        ROOT / "红方规则对抗MARL场景需求说明.docx",
        "红方规则对抗 MARL 场景需求说明",
    )
    build_doc(
        ROOT / "红方规则对抗开发与验收计划.md",
        ROOT / "红方规则对抗开发与验收计划.docx",
        "红方规则对抗开发与验收计划",
    )


if __name__ == "__main__":
    main()
